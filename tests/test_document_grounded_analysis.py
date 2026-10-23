"""Actual-document authority with native scoring; isolated, zero-cost receipts."""
from copy import deepcopy
from contextlib import closing
import ast
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, Mock

from database import db_manager, analysis_cache_manager as cache
from tailoring import document_grounded_analysis as document
from tailoring.initial_analysis_explainer import load_application_initial, load_frozen_document_report
from tailoring.phase9e_blueprint_selection import (
    build_original_resume_starting_snapshot, evaluate_starting_snapshot,
    fingerprint_value, Phase9EDecisionError, _build_phase9e_starting_analysis,
)
from analysis_stability.stable_evidence_scoring import canonicalise_requirements
from analysis_stability.requirement_context_policy import current_scoring_version

ROOT = Path(__file__).resolve().parents[1]
JD = 'Experience in API development using FastAPI'
BULLET = 'Developed REST API endpoints using FastAPI for authentication and data retrieval.'
PROFILE = {'projects': [{'title': 'API service', 'bullets': [BULLET]}], 'skills': {}}
EMPTY_TEXT = 'Built accessible web pages and managed presentation layouts. '


def report(text=EMPTY_TEXT):
    result = document.analyze_actual_resume_against_jd(resume_text=text, jd_text=JD,
        jd_profile={'requirements': [JD]}, resume_profile=PROFILE)
    result['meta'] = {'document_analysis': result.pop('document_analysis')}
    return result


def exact_jd():
    canonical = canonicalise_requirements(jd_profile={'requirements': [JD]}, raw_jd_text=JD)
    return {'raw_text': JD, 'jd_profile': {'requirements': [JD]}, 'canonicalisation': canonical}


class DocumentInputTests(unittest.TestCase):
    def test_real_analyze_entry_uses_document_authority_with_fake_provider_results(self):
        from datetime import datetime
        from analyzer import compute_overall_score
        tree = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        function = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'run_resume_analysis')
        namespace = {'st': Mock(), 'datetime': datetime, 'deepcopy': deepcopy,
            'get_active_model': lambda _: 'fake', 'ATS_PASS_THRESHOLD': 60,
            'compute_overall_score': compute_overall_score, 'summarise_overall': lambda _: 'fake summary',
            'extract_resume_profile': lambda _: deepcopy(PROFILE),
            'extract_jd_profile': lambda _: {'requirements': [JD]},
            'analyse_keyword_match': lambda *args: {'present': [{'requirement': JD, 'match_type': 'direct'}], 'keyword_match_score': 100},
            'analyse_bullets': lambda *args: {'bullet_quality_avg': 100},
            'analyse_jargon': lambda *args: {'jargon_score': 100},
            'analyse_structure': lambda *args, **kwargs: {'structure_score': 100},
            'analyse_degree_alignment': lambda *args: {'degree_alignment_score': 100}}
        exec(compile(ast.Module(body=[function], type_ignores=[]), str(ROOT / 'app.py'), 'exec'), namespace)
        current = namespace['run_resume_analysis'](EMPTY_TEXT, JD, 'Computing')
        document.validate_document_report(current)
        self.assertEqual(current['stable_analysis']['scoring_version'], current_scoring_version())
        self.assertEqual(current['stable_analysis']['canonical_requirements'][0]['match_label'], 'none')
        self.assertEqual(current['keyword_match']['keyword_match_score'], 100)
        self.assertEqual(current['overall_score'], 100)  # Existing legacy quality arithmetic only.

    def test_absent_profile_evidence_has_no_document_credit(self):
        current = report()
        self.assertEqual(current['stable_analysis']['canonical_requirements'][0]['match_label'], 'none')
        self.assertEqual(current['resume_profile']['projects'], [])
        self.assertEqual(current['raw_resume_text'], EMPTY_TEXT)
        self.assertEqual(current['stable_analysis']['scoring_version'], current_scoring_version())

    def test_opportunity_is_native_grounded_and_does_not_change_score(self):
        current = report()
        before = deepcopy(current)
        rows = document.build_tailoring_evidence_opportunities(current,
            [{'id': 1, 'title': 'API service', 'description': BULLET}])
        self.assertEqual(rows[0]['state'], 'GROUNDED_EVIDENCE_AVAILABLE')
        self.assertEqual(rows[0]['additional_grounded_evidence'][0]['native_label'], 'direct')
        self.assertEqual(current, before)
        self.assertNotIn('projected_score', rows[0])

    def test_represented_document_accepts_native_evidence_without_id_drift(self):
        initial, final = report(), report(BULLET)
        a, b = [r['stable_analysis']['canonical_requirements'][0] for r in (initial, final)]
        self.assertEqual(a['requirement_id'], b['requirement_id'])
        self.assertEqual(b['match_label'], 'direct')
        self.assertEqual(document.build_tailoring_evidence_opportunities(final, [])[0]['state'], 'ALREADY_REPRESENTED')

    def test_same_document_identity_has_same_complete_native_receipt(self):
        self.assertEqual(report(BULLET), report(BULLET))
        self.assertNotEqual(report()['stable_analysis'], report(BULLET)['stable_analysis'])

    def test_literal_source_filter_has_token_boundaries(self):
        self.assertEqual(document.represented_profile({'skills': {'Languages': ['C', 'Python']}}, 'Created CPython widgets')['skills'], {'Languages': []})

    def test_contact_and_layout_metadata_survive_evidence_source_bounding(self):
        profile = {**PROFILE, 'personal_info': {'name': 'Example Candidate'}, 'layout': {'order': ['projects']}}
        bounded = document.represented_profile(profile, EMPTY_TEXT)
        self.assertEqual(bounded['personal_info'], profile['personal_info'])
        self.assertEqual(bounded['layout'], profile['layout'])
        self.assertEqual(bounded['projects'], [])

    def test_bad_authority_and_scope_fail_closed(self):
        current = report()
        current['raw_resume_text'] = BULLET
        with self.assertRaisesRegex(ValueError, 'provenance'):
            document.validate_document_report(current)
        with self.assertRaisesRegex(ValueError, 'canonical'):
            document.analyze_actual_resume_against_jd(resume_text=BULLET, jd_text=JD,
                jd_profile={'requirements': [JD]}, canonical_requirements=[])

    def test_application_source_reuses_frozen_receipt_without_legacy_rescore(self):
        current = report(BULLET)
        current['frozen_initial_resume_report'] = deepcopy(current)
        snapshot = build_original_resume_starting_snapshot(current)
        with patch('tailoring.phase9e_blueprint_selection.build_stable_analysis', side_effect=AssertionError('No legacy rescore')):
            stable, _, _ = _build_phase9e_starting_analysis(starting_snapshot=snapshot, jd=exact_jd(), preferred_requirements=[])
        self.assertEqual(stable, current['stable_analysis'])
        self.assertEqual(snapshot['analysis_source'], 'Frozen Initial Résumé Analysis')

    def test_master_document_is_analyzed_independently(self):
        snapshot = {'resume_profile_snapshot': PROFILE, 'resume_text_snapshot': BULLET, 'document_analysis_mode': True}
        with patch('tailoring.phase9e_blueprint_selection.build_stable_analysis', side_effect=AssertionError('No legacy rescore')):
            stable, _, _ = _build_phase9e_starting_analysis(starting_snapshot=snapshot, jd=exact_jd(), preferred_requirements=[])
        self.assertEqual(stable, report(BULLET)['stable_analysis'])
        self.assertNotEqual(stable, report()['stable_analysis'])

    def test_same_identity_inconsistent_receipt_fails_closed(self):
        current = report(BULLET)
        frozen = deepcopy(current)
        frozen['stable_analysis']['deterministic_alignment_score'] += 1
        document.seal_document_report(frozen)
        current['frozen_initial_resume_report'] = frozen
        with self.assertRaisesRegex(Phase9EDecisionError, 'inconsistent'):
            build_original_resume_starting_snapshot(current)

    def test_comparison_cache_retains_document_authority_not_usage_metadata(self):
        from tailoring.phase9e_blueprint_selection_ui import _ranking_report_snapshot
        current = report()
        self.assertEqual(_ranking_report_snapshot(current), current)
        first = build_original_resume_starting_snapshot(current)
        current['meta']['api_cost_summary'] = {'total': 123}
        second = build_original_resume_starting_snapshot(current)
        self.assertEqual(first['starting_snapshot_fingerprint'], second['starting_snapshot_fingerprint'])

    def test_passive_document_currentness_never_rebuilds(self):
        from tailoring.jd_user_input_overrides import refresh_application_session_analysis_report
        current = report()
        with patch('tailoring.jd_user_input_overrides.build_stable_analysis', side_effect=AssertionError('No passive scoring')):
            view = refresh_application_session_analysis_report(current)
        self.assertEqual(view['stable_analysis'], current['stable_analysis'])
        self.assertFalse(view['meta']['stable_analysis_currentness']['rebuilt'])

    def test_stale_document_opportunities_fail_closed_without_rescore(self):
        current = report()
        with patch.object(document, 'current_ab_identity', return_value={'scoring_version': 'future'}):
            with self.assertRaisesRegex(ValueError, 'stale'):
                document.build_tailoring_evidence_opportunities(current, [])

    def test_results_order_and_one_taxonomy_surface(self):
        source = (ROOT / 'app.py').read_text(encoding='utf-8')
        ast.parse(source)
        index = source.index('    if report:\n        overall_score')
        section = source[index:source.index('        phase9a_jd_record:', index)]
        self.assertLess(source.index('st.header("Application workflow")'), source.index('render_application_bound_jd(st'))
        self.assertNotIn('st.header("Results")', source)
        self.assertLess(section.index('render_application_workflow_overview('), section.index('render_application_initial_analysis(st'))
        self.assertLess(section.index('render_application_initial_analysis(st'), section.index('render_tailoring_evidence_opportunities(st'))
        self.assertNotIn('render_jd_taxonomy_status(', source)


class DocumentReceiptTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.path = Path(temp.name) / 'applications.db'
        patcher = patch.object(db_manager, 'DB_PATH', self.path)
        patcher.start(); self.addCleanup(patcher.stop)
        db_manager.init_db(); cache.init_analysis_cache()
        self.app = db_manager.create_empty_application_session()
        for name in ('litellm.completion', 'litellm.embedding', 'socket.create_connection'):
            guard = patch(name, side_effect=AssertionError('No live providers'))
            guard.start(); self.addCleanup(guard.stop)

    def persist(self, current, key='doc-input', reused=None):
        return cache.persist_explicit_application_analysis(application_id=self.app,
            input_fingerprint=key, report=current, analysis_model='fake', resume_filename='resume.docx',
            raw_jd_text=JD, reused_snapshot=reused)

    def test_first_document_receipt_is_frozen_and_exact_cache_reused(self):
        current = report()
        saved = cache.save_analysis_snapshot(application_id=self.app, input_fingerprint='doc-input',
            report=current, analysis_model='fake', resume_filename='resume.docx')
        with self.assertRaisesRegex(ValueError, 'No authoritative'):
            load_application_initial(self.app, resume_only=True)
        bound = self.persist(current, reused=saved)
        self.assertEqual(bound['analysis_id'], saved['analysis_id'])
        initial = load_application_initial(self.app, resume_only=True)
        self.assertEqual(initial['provenance'], 'EXACT_SAVED_REUSE')
        self.persist(report(BULLET), key='new-doc')
        self.assertEqual(load_application_initial(self.app, resume_only=True), initial)
        self.assertEqual(load_frozen_document_report(self.app)['stable_analysis'], current['stable_analysis'])
        with self.assertRaisesRegex(ValueError, 'Frozen'):
            cache.save_analysis_snapshot(application_id=self.app, input_fingerprint='changed', report=report(BULLET),
                analysis_model='fake', resume_filename='resume.docx', analysis_id=saved['analysis_id'])

    def test_discovery_history_is_not_automatically_document_initial(self):
        historical = report()
        historical['meta'].pop('document_analysis')
        saved = self.persist(historical, key='discovery')
        binding = {'initial_analysis_id': saved['analysis_id'], 'raw_jd_text': JD,
            'snapshot': {'stable_analysis': historical['stable_analysis'], 'id': 11}, 'discovered_job_id': 88}
        binding['binding_fingerprint'] = fingerprint_value(binding)
        with closing(db_manager._connect()) as conn:
            conn.execute('UPDATE applications SET entry_binding_json=? WHERE id=?', (json.dumps(binding), self.app)); conn.commit()
        with self.assertRaisesRegex(ValueError, 'No authoritative'):
            load_application_initial(self.app, resume_only=True)
        self.assertEqual(load_application_initial(self.app)['provenance'], 'JOB_FINDER_INITIAL_MATCH')
        self.persist(report(BULLET))
        self.assertEqual(load_application_initial(self.app)['provenance'], 'JOB_FINDER_INITIAL_MATCH')
        self.assertEqual(load_application_initial(self.app, resume_only=True)['stable_analysis'], report(BULLET)['stable_analysis'])

    def test_exact_document_receipt_with_discovery_history_requires_explicit_binding(self):
        current = report()
        saved = cache.save_analysis_snapshot(application_id=self.app, input_fingerprint='doc-input',
            report=current, analysis_model='fake', resume_filename='resume.docx')
        with self.assertRaisesRegex(ValueError, 'No authoritative'):
            load_application_initial(self.app, resume_only=True)
        self.persist(current, reused=saved)
        self.assertEqual(load_application_initial(self.app, resume_only=True)['identity']['initial_analysis_id'], saved['analysis_id'])

    def test_passive_ui_has_no_writes_models_or_projection(self):
        from streamlit.testing.v1 import AppTest
        # Resolve the existing optional review UI dependencies outside the
        # script-run deadline; all rendering and no-write assertions stay active.
        from tailoring import initial_analysis_explainer_ui, unresolved_requirement_review_ui
        self.persist(report())
        before = self.path.read_bytes()
        app = AppTest.from_string(f'''import streamlit as st
from tailoring.initial_analysis_explainer_ui import render_application_initial_analysis, render_tailoring_evidence_opportunities
from tailoring.initial_analysis_explainer import load_frozen_document_report
st.header('Results')
render_application_initial_analysis(st, {self.app}, resume_only=True, show_comparison=False)
render_tailoring_evidence_opportunities(st, load_frozen_document_report({self.app}))
''').run(timeout=30)
        self.assertFalse(app.exception)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertIn('Initial Résumé Analysis', [e.label for e in app.expander])
        self.assertIn('Evidence available for tailoring', [e.label for e in app.expander])


class FinalDocumentTests(unittest.TestCase):
    def generation(self, path):
        from tests.test_phase8_scoring_correctness import _generation_for_mapping
        row = report()['stable_analysis']['canonical_requirements'][0]
        generation = _generation_for_mapping(requirement_id=row['requirement_id'], requirement_text=JD, bullet=BULLET)
        generation['fit_result']['docx_path'] = str(path)
        return generation

    def test_phase8_actual_document_presence_and_missing_artifact_fail_closed(self):
        from docx import Document
        from tailoring.phase8_verification import build_phase8_verification
        from tailoring.phase8_auto_verification import phase8_verification_is_current
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'final.docx'
            doc = Document(); doc.add_paragraph(EMPTY_TEXT * 5); doc.save(path)
            generation = self.generation(path)
            missing = build_phase8_verification(baseline_report=report(), generation_state=generation, raw_jd_text=JD)
            self.assertEqual(missing['after_stable_analysis']['canonical_requirements'][0]['match_label'], 'none')
            doc.add_paragraph(BULLET); doc.save(path)
            self.assertFalse(phase8_verification_is_current(missing, generation))
            final = build_phase8_verification(baseline_report=report(), generation_state=generation, raw_jd_text=JD)
            self.assertEqual(final['before_stable_analysis'], report()['stable_analysis'])
            self.assertEqual(final['after_stable_analysis']['canonical_requirements'][0]['match_label'], 'direct')
            self.assertTrue(final['comparison']['canonical_requirement_ids_stable'])
            self.assertTrue(final['document_analysis']['final_document_sha256'])
            self.assertEqual(final['comparison']['improved_requirements'][0]['delta_reason'], 'NEW_FINAL_RESUME_EVIDENCE')
            from database.tailoring_verification_manager import get_tailoring_verification_approval_gate
            from tailoring.phase8_verification import refresh_phase8_readiness
            from tailoring.phase8_auto_verification import phase8_verification_refresh_reason
            with patch('database.tailoring_generation_control.get_tailoring_generation', return_value=generation), \
                 patch('database.tailoring_verification_manager.get_latest_tailoring_verification', return_value=final):
                doc.add_paragraph('Changed actual final artifact'); doc.save(path)
                gate = get_tailoring_verification_approval_gate(1, 'generation')
                self.assertIn('actual_final_document_changed_or_unavailable', gate['reasons'])
                self.assertEqual(phase8_verification_refresh_reason(final, generation), 'actual_final_document_changed_or_unavailable')
                with self.assertRaisesRegex(ValueError, 'document changed'):
                    refresh_phase8_readiness(final, generation)
            generation['fit_result']['docx_path'] = str(Path(temp) / 'absent.docx')
            with self.assertRaises(ValueError):
                build_phase8_verification(baseline_report=report(), generation_state=generation, raw_jd_text=JD)


class DocumentSelectionTests(unittest.TestCase):
    def test_document_decision_integrity_and_legacy_blueprint_exclusion(self):
        from contextlib import ExitStack
        from database import jd_library_manager, tailoring_version_manager, global_blueprint_manager
        import database.phase9f_exact_verified_reuse_manager as reuse
        from tests.phase9e_test_support import seed_phase9e_database
        from tailoring.phase9e_blueprint_selection import build_phase9e_decision, verify_decision_integrity, build_effective_tailoring_report
        with tempfile.TemporaryDirectory() as temp, ExitStack() as stack:
            path = Path(temp) / 'selection.db'
            for module, attribute in ((db_manager, 'DB_PATH'), (jd_library_manager, 'DB_PATH'),
                    (tailoring_version_manager, 'DB_PATH'), (global_blueprint_manager, 'BLUEPRINT_ARTIFACT_ROOT'),
                    (reuse, 'BLUEPRINT_ARTIFACT_ROOT')):
                stack.enter_context(patch.object(module, attribute, getattr(module, attribute)))
            state = seed_phase9e_database(path, different_original=True)
            jd = jd_library_manager.get_exact_job_description_for_application(94)
            native = document.analyze_actual_resume_against_jd(resume_text=EMPTY_TEXT,
                jd_text=jd['raw_text'], jd_profile=jd['jd_profile'], resume_profile=PROFILE)
            native['meta'] = {'document_analysis': native.pop('document_analysis')}
            native['frozen_initial_resume_report'] = deepcopy(native)
            decision = build_phase9e_decision(application_id=94, application_report=native,
                exact_jd=jd, active_blueprints=[state['blueprint']], selected_source='original_resume')
            verify_decision_integrity(decision)
            self.assertEqual(decision['comparison']['analysis_source'], 'Frozen Initial Résumé Analysis')
            self.assertEqual(decision['comparison']['stable_analysis_snapshot'], native['stable_analysis'])
            effective = build_effective_tailoring_report(native, decision)
            document.validate_document_report(effective)
            self.assertEqual(effective['frozen_initial_resume_report']['stable_analysis'], native['stable_analysis'])
            with self.assertRaisesRegex(Phase9EDecisionError, 'not active'):
                build_phase9e_decision(application_id=94, application_report=native,
                    exact_jd=jd, active_blueprints=[state['blueprint']], selected_source='global_blueprint',
                    selected_blueprint_id=state['blueprint']['blueprint_id'])
