from __future__ import annotations

import unittest
from io import BytesIO

import openpyxl

from tailoring.application_tracker_export import (
    build_application_tracker_workbook,
)


ROWS = [
    {
        "tracked_job_id": 1,
        "application_id": 10,
        "company": "Example Co",
        "job_title": "Software Engineer",
        "location": "Singapore",
        "overall_score": 82,
        "applied": 1,
        "applied_at": "2026-09-01",
        "status": "interview",
        "completed": 0,
        "completed_at": None,
        "notes": "Phone screen",
        "job_url": "https://example.com/job",
        "source_type": "application_session",
    },
    {
        "tracked_job_id": 2,
        "application_id": None,
        "company": "Manual Co",
        "job_title": "Data Engineer",
        "location": "",
        "overall_score": None,
        "applied": 0,
        "applied_at": None,
        "status": "not_applied",
        "completed": 0,
        "completed_at": None,
        "notes": "",
        "job_url": "",
        "source_type": "manual",
    },
]


class ApplicationTrackerExportTests(unittest.TestCase):
    def test_workbook_contains_expected_sheets(self):
        payload = build_application_tracker_workbook(ROWS, [])
        workbook = openpyxl.load_workbook(BytesIO(payload))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Applications",
                "Summary",
                "Weekly Trend",
                "Weekly Cohort",
                "Status History",
            ],
        )

    def test_applications_sheet_contains_human_status_label(self):
        payload = build_application_tracker_workbook(ROWS, [])
        workbook = openpyxl.load_workbook(BytesIO(payload))
        worksheet = workbook["Applications"]
        headers = [cell.value for cell in worksheet[1]]
        status_column = headers.index("Status") + 1
        self.assertEqual(worksheet.cell(2, status_column).value, "Interview")

    def test_summary_counts_applied_jobs(self):
        payload = build_application_tracker_workbook(ROWS, [])
        workbook = openpyxl.load_workbook(BytesIO(payload), data_only=True)
        worksheet = workbook["Summary"]
        values = {
            worksheet.cell(row, 1).value: worksheet.cell(row, 2).value
            for row in range(2, worksheet.max_row + 1)
        }
        self.assertEqual(values["Tracked Jobs"], 2)
        self.assertEqual(values["Applied"], 1)


if __name__ == "__main__":
    unittest.main()
