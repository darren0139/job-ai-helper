"""Three document roles, using native receipts and existing source currentness."""
from contextlib import ExitStack
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from tailoring import document_grounded_analysis as document
from tests.test_document_grounded_analysis import report, JD, BULLET, EMPTY_TEXT, PROFILE
from tests import test_document_grounded_analysis as support

ROOT = Path(__file__).resolve().parents[1]


def selected_report():
    selected = report(BULLET)
    selected['frozen_initial_resume_report'] = report()
    selected['meta']['phase9e_starting_context'] = {'selected_source': 'base_resume',
        'decision_id': 'master-decision', 'starting_snapshot_fingerprint': 'master-snapshot'}
    return selected


class WorkflowHierarchyTests(unittest.TestCase):
    def test_selected_receipt_is_native_and_initial_is_immutable(self):
        current = selected_report()
        before = deepcopy(current)
        with patch.object(document, 'analyze_actual_resume_against_jd', side_effect=AssertionError('No rescore')):
            view = document.selected_base_analysis(current)
        self.assertFalse(view['same_as_initial'])
        self.assertEqual(view['stable_analysis'], report(BULLET)['stable_analysis'])
        self.assertNotEqual(view['stable_analysis'], current['frozen_initial_resume_report']['stable_analysis'])
        self.assertEqual(current, before)
        self.assertEqual(view['starting_context']['selected_source'], 'base_resume')

    def test_switch_back_reuses_exact_initial(self):
        master = selected_report()
        original = deepcopy(master['frozen_initial_resume_report'])
        original['frozen_initial_resume_report'] = deepcopy(original)
        original['meta']['phase9e_starting_context'] = {'selected_source': 'original_resume'}
        self.assertFalse(document.selected_base_analysis(master)['same_as_initial'])
        view = document.selected_base_analysis(original)
        self.assertTrue(view['same_as_initial'])
        self.assertEqual(view['stable_analysis'], master['frozen_initial_resume_report']['stable_analysis'])

    def test_same_document_conflicting_outcomes_fail_closed(self):
        current = report()
        current['frozen_initial_resume_report'] = deepcopy(current)
        current['stable_analysis']['deterministic_alignment_score'] += 1
        document.seal_document_report(current)
        with self.assertRaisesRegex(ValueError, 'conflicting'):
            document.selected_base_analysis(current)

    def test_incompatible_selected_identity_fails_closed(self):
        current = selected_report()
        current['raw_jd_text'] += ' Experience in SQL'
        document.seal_document_report(current)
        with self.assertRaisesRegex(ValueError, 'incompatible'):
            document.selected_base_analysis(current)

    def verification(self, current):
        from docx import Document
        from tailoring.phase8_verification import build_phase8_verification
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        path = Path(temp.name) / 'final.docx'
        doc = Document(); doc.add_paragraph(EMPTY_TEXT * 5); doc.add_paragraph(BULLET); doc.save(path)
        generation = support.FinalDocumentTests().generation(path)
        return build_phase8_verification(baseline_report=current, generation_state=generation, raw_jd_text=JD), generation

    def test_two_comparisons_join_native_receipts_without_replacing_initial(self):
        current = selected_report()
        before = deepcopy(current)
        result, _ = self.verification(current)
        initial_ab = result['job_match_ab']
        selected_ab = result['selected_base_job_match_ab']
        self.assertTrue(initial_ab['comparable']); self.assertTrue(selected_ab['comparable'])
        self.assertEqual(result['before_stable_analysis'], current['frozen_initial_resume_report']['stable_analysis'])
        self.assertEqual(result['selected_base_analysis']['stable_analysis'], current['stable_analysis'])
        self.assertEqual(initial_ab['initial']['score'], current['frozen_initial_resume_report']['stable_analysis']['deterministic_alignment_score'])
        self.assertEqual(selected_ab['initial']['score'], current['stable_analysis']['deterministic_alignment_score'])
        self.assertEqual(initial_ab['tailored'], selected_ab['tailored'])
        self.assertNotEqual(initial_ab['summary']['score_delta'], selected_ab['summary']['score_delta'])
        self.assertEqual(result['after_stable_analysis']['canonical_requirements'][0]['match_label'], 'direct')
        self.assertEqual(current, before)

    def test_same_source_collapses_comparison_and_readiness_keeps_receipts(self):
        from tailoring.phase8_verification import refresh_phase8_readiness
        result, generation = self.verification(report())
        self.assertTrue(result['selected_base_analysis']['same_as_initial'])
        self.assertIsNone(result['selected_base_job_match_ab'])
        refreshed = refresh_phase8_readiness(result, generation)
        for key in ('job_match_ab', 'selected_base_analysis', 'selected_base_job_match_ab', 'document_analysis'):
            self.assertEqual(refreshed[key], result[key])

    def test_selected_receipts_are_immutable_persistence_fields(self):
        from database.tailoring_verification_manager import _IMMUTABLE_ANALYSIS_FIELDS
        self.assertIn('selected_base_analysis', _IMMUTABLE_ANALYSIS_FIELDS)
        self.assertIn('selected_base_job_match_ab', _IMMUTABLE_ANALYSIS_FIELDS)

    def test_ui_umbrella_and_order_before_explicit_execution(self):
        source = (ROOT / 'app.py').read_text(encoding='utf-8')
        section = source[source.index('elif page == "Application Sessions":'):source.index('        phase9e_ready =')]
        self.assertNotIn('st.header("Results")', section)
        ordered = ['st.header("Application workflow")', 'render_application_bound_jd(st',
            'render_application_workflow_overview(', 'Step 1 — Initial Résumé Analysis',
            'Step 2 — Evidence available for tailoring', 'Step 3 — Tailoring base decision',
            'Step 4 — Selected base', 'render_selected_base_analysis(st', 'Step 5 — Tailor résumé',
            'render_phase9f_reuse_execution(']
        # The no-current-report branch also renders Initial; inspect its later
        # occurrence for the normal established workflow.
        established = section.index('    if report:\n        overall_score')
        positions = [section.index(token, established if token.startswith('Step 1') else 0) for token in ordered]
        self.assertEqual(positions, sorted(positions))

    def test_passive_selected_ui_collapses_same_and_shows_different(self):
        from streamlit.testing.v1 import AppTest
        from tailoring import initial_analysis_explainer_ui as ui
        for current, expect_panel in ((report(), False), (selected_report(), True)):
            before = deepcopy(current)
            with patch.object(document, 'analyze_actual_resume_against_jd', side_effect=AssertionError('No scoring')), \
                    patch.object(ui, 'render_jd_taxonomy_status', side_effect=AssertionError('No duplicate taxonomy')):
                app = AppTest.from_string('''import streamlit as st
from tailoring.initial_analysis_explainer_ui import render_selected_base_analysis
render_selected_base_analysis(st, st.session_state['receipt'])
''')
                app.session_state['receipt'] = current
                app.run(timeout=30)
            self.assertFalse(app.exception)
            self.assertEqual('Selected Base Analysis' in [e.label for e in app.expander], expect_panel)
            if not expect_panel:
                self.assertFalse(app.metric)
                self.assertTrue(any('Using frozen Initial' in c.value for c in app.caption))
            self.assertEqual(current, before)

    def test_ui_labels_two_deltas_and_same_source_single_comparison(self):
        from streamlit.testing.v1 import AppTest
        for current, second in ((report(), False), (selected_report(), True)):
            result, _ = self.verification(current)
            app = AppTest.from_string('''import streamlit as st
from tailoring.phase8_verification_ui import _render_result
_render_result(st.session_state['result'])
''')
            app.session_state['result'] = result
            app.run(timeout=30)
            self.assertFalse(app.exception)
            self.assertIn('Initial Job Match', [m.label for m in app.metric])
            self.assertIn('Final Job Match', [m.label for m in app.metric])
            text = [m.value for m in app.markdown]
            self.assertIn('#### Initial vs Final', text)
            self.assertEqual('#### Selected Base vs Final' in text, second)
            self.assertEqual(text.count('### Step 6 — Final Résumé Analysis / Phase8'), 1)
            self.assertLess(text.index('### Step 6 — Final Résumé Analysis / Phase8'), text.index('### Step 7 — Analysis comparison'))
            self.assertFalse(any('Baseline before fitting' in value for value in text))

    def test_17_24_31_presentation_keeps_session_and_tailoring_deltas_distinct(self):
        # A presentation fixture, not fabricated native scoring evidence. The
        # real-document test above separately proves native receipt parity.
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_string('''from tailoring.phase8_verification_ui import _render_document_comparison
_render_document_comparison({'comparable': True, 'requirements': [],
    'summary': {'initial_score': 17, 'tailored_score': 31, 'score_delta': 14}}, 'Initial vs Final')
_render_document_comparison({'comparable': True, 'requirements': [],
    'summary': {'initial_score': 24, 'tailored_score': 31, 'score_delta': 7}}, 'Selected Base vs Final')
''').run(timeout=30)
        self.assertFalse(app.exception)
        self.assertEqual([(m.label, m.value) for m in app.metric],
            [('Initial', '17'), ('Final', '31'), ('Job Match Δ', '14'),
             ('Selected Base', '24'), ('Final', '31'), ('Job Match Δ', '7')])


class NativeSourceDecisionTests(unittest.TestCase):
    def test_master_selection_and_switch_back_preserve_saved_initial_and_invalidate_source(self):
        from database import db_manager, jd_library_manager, tailoring_version_manager, global_blueprint_manager
        from database import analysis_cache_manager as cache
        import database.phase9f_exact_verified_reuse_manager as reuse
        from tests.phase9e_test_support import seed_phase9e_database
        from tailoring.phase9e_blueprint_selection import (build_base_resume_starting_snapshot,
            build_phase9e_decision, build_effective_tailoring_report, generation_binding_identity)
        from tailoring.tailoring_generation_fingerprint import generation_matches_phase9e_binding
        from tailoring.initial_analysis_explainer import load_application_initial
        with tempfile.TemporaryDirectory() as temp, ExitStack() as stack:
            path = Path(temp) / 'source.db'
            for module, attribute in ((db_manager, 'DB_PATH'), (jd_library_manager, 'DB_PATH'),
                    (tailoring_version_manager, 'DB_PATH'), (global_blueprint_manager, 'BLUEPRINT_ARTIFACT_ROOT'),
                    (reuse, 'BLUEPRINT_ARTIFACT_ROOT')):
                stack.enter_context(patch.object(module, attribute, getattr(module, attribute)))
            state = seed_phase9e_database(path, different_original=True)
            jd = jd_library_manager.get_exact_job_description_for_application(94)
            initial = document.analyze_actual_resume_against_jd(resume_text=EMPTY_TEXT,
                jd_text=jd['raw_text'], jd_profile=jd['jd_profile'], resume_profile=PROFILE)
            initial['meta'] = {'document_analysis': initial.pop('document_analysis')}
            cache.persist_explicit_application_analysis(application_id=94, input_fingerprint='original-doc',
                report=initial, analysis_model='fake', resume_filename='original.docx', raw_jd_text=jd['raw_text'])
            frozen = load_application_initial(94, resume_only=True)
            database_before = path.read_bytes()
            initial['frozen_initial_resume_report'] = deepcopy(initial)
            snapshot = build_base_resume_starting_snapshot(state['base_resume'], state['base_resume_artifact'])
            decisions = {}
            for source in ('original_resume', 'base_resume'):
                decisions[source] = build_phase9e_decision(application_id=94, application_report=initial,
                    exact_jd=jd, active_blueprints=[], selected_source=source,
                    base_resume_starting_snapshot=snapshot, selection_mode='manual')
            original = build_effective_tailoring_report(initial, decisions['original_resume'])
            master = build_effective_tailoring_report(initial, decisions['base_resume'])
            self.assertTrue(document.selected_base_analysis(original)['same_as_initial'])
            self.assertFalse(document.selected_base_analysis(master)['same_as_initial'])
            self.assertEqual(original['stable_analysis'], frozen['stable_analysis'])
            self.assertEqual(master['raw_resume_text'], snapshot['resume_text_snapshot'])
            self.assertEqual(master['stable_analysis'], decisions['base_resume']['comparison']['stable_analysis_snapshot'])
            generation = {'generation_settings': {'phase9e_binding': generation_binding_identity(decisions['original_resume'])}}
            self.assertTrue(generation_matches_phase9e_binding(generation, generation_binding_identity(decisions['original_resume'])))
            self.assertFalse(generation_matches_phase9e_binding(generation, generation_binding_identity(decisions['base_resume'])))
            self.assertEqual(load_application_initial(94, resume_only=True), frozen)
            self.assertEqual(path.read_bytes(), database_before)
